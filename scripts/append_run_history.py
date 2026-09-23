"""Append one row to RunHistory.xlsx. Called from the scenario .jl scripts as a subprocess."""
import argparse
import sys

import openpyxl
from openpyxl.styles import Font, Alignment

FONT_NAME = "Arial"


def parse_args():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--xlsx", required=True)
    p.add_argument("--timestamp", required=True)
    p.add_argument("--scenario", required=True)
    p.add_argument("--runner-script", required=True)
    p.add_argument("--csv", required=True)
    p.add_argument("--mp4", required=True)
    p.add_argument("--hist-mp4", default="")
    p.add_argument("--pathfinding-time", required=True, type=float)
    return p.parse_args()


def main():
    args = parse_args()

    try:
        wb = openpyxl.load_workbook(args.xlsx)
    except PermissionError:
        print(
            f"ΣΦΑΛΜΑ: το {args.xlsx} είναι ανοιχτό (πιθανόν σε Excel). "
            "Κλείσ' το και ξανατρέξε το script.",
            file=sys.stderr,
        )
        sys.exit(1)

    ws = wb["Runs"]
    table = ws.tables["RunsTable"]
    first_cell, last_cell = table.ref.split(":")
    last_row = int("".join(ch for ch in last_cell if ch.isdigit()))
    new_row = last_row + 1

    values = {
        1: "=ROW()-1",
        2: args.timestamp,
        3: args.scenario,
        4: args.runner_script,
        5: args.csv,
        6: args.mp4,
        7: args.hist_mp4 or None,
        8: round(args.pathfinding_time, 6),
        9: None,  # Αλλαγές Κώδικα Πριν την Εκτέλεση — χειροκίνητη συμπλήρωση
    }
    for col_idx, value in values.items():
        cell = ws.cell(row=new_row, column=col_idx, value=value)
        cell.font = Font(name=FONT_NAME)
        if col_idx == 9:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=new_row, column=8).number_format = "0.000"
    ws.row_dimensions[new_row].height = 30

    last_col = "".join(ch for ch in last_cell if ch.isalpha())
    table.ref = f"{first_cell}:{last_col}{new_row}"

    try:
        wb.save(args.xlsx)
    except PermissionError:
        print(
            f"ΣΦΑΛΜΑ: αποτυχία αποθήκευσης του {args.xlsx} (πιθανόν ανοιχτό σε Excel). "
            "Κλείσ' το και ξανατρέξε το script.",
            file=sys.stderr,
        )
        sys.exit(1)

    print(f"RunHistory.xlsx: προστέθηκε γραμμή {new_row - 1} (Run ID {new_row - 1}).")


if __name__ == "__main__":
    main()
